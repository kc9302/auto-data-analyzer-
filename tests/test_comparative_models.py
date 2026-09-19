"""
Tests for Comparative Benchmark Models & Extended Leaderboard Excel Reporting.
Verifies that:
1. MLScoutEngine calculates Global Stat, Demographic Segment Rule baselines, and individual Lift % for all models.
2. Segment slices (age-based or quantile-based) are correctly evaluated.
3. ExcelReportBuilder renders 10-column leaderboard and slice benchmark tables in Sheet 4.
"""
import os
import sys
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import pytest
import numpy as np
import pandas as pd
import openpyxl
from src.ml_scout.engine import MLScoutEngine
from src.presenter.excel_builder import ExcelReportBuilder


def test_comparative_models_and_lift_calculation():
    """Verifies that each model in leaderboard has comparative lift columns and roles."""
    np.random.seed(42)
    n = 120
    # Create dataset with an explicit 'age' column to trigger demographic slice analysis
    ages = np.random.randint(20, 65, size=n)
    salary = np.random.normal(5000, 1000, size=n)
    experience = np.random.uniform(1, 15, size=n)
    y = ((salary > 5200) & (experience > 5)).astype(int)

    df_X = pd.DataFrame({"age": ages, "salary": salary, "experience": experience})
    s_y = pd.Series(y, name="target")

    scout = MLScoutEngine(random_seed=42, cv_folds=3)
    res = scout.run_scout(df_X, s_y)

    assert "lift_analysis" in res
    lift = res["lift_analysis"]
    assert "champion_model" in lift
    assert "global_baseline_score" in lift
    assert "segment_baseline_score" in lift
    assert "lift_vs_global_pct" in lift
    assert "lift_vs_segment_pct" in lift
    assert "segment_slices" in lift

    # Verify segment slices detected 'age'
    slices = lift["segment_slices"]
    assert len(slices) >= 2
    assert any("청년층" in s["segment_name"] or "중년층" in s["segment_name"] or "장년" in s["segment_name"] for s in slices)
    for s in slices:
        assert "sample_count" in s
        assert "sample_share_pct" in s
        assert "baseline_score" in s
        assert "champion_score" in s
        assert "lift_pct" in s
        assert "interpretation" in s

    # Verify each leaderboard item has enriched comparative fields
    board = res["leaderboard"]
    for m in board:
        assert "lift_vs_global_pct" in m
        assert "lift_vs_segment_pct" in m
        assert "model_role" in m
        assert "algorithm_family" in m
        assert "complexity" in m
        assert "verdict_rationale" in m
        assert len(m["verdict_rationale"]) > 5

    # Check champion model role
    champ_row = next(m for m in board if m["model"] == res["best_model"])
    assert "🏆 챔피언" in champ_row["model_role"]


def test_excel_leaderboard_sheet_structure(tmp_path):
    """Verifies that Sheet 4 (AutoML_리더보드) contains 3 sections and 10 columns."""
    np.random.seed(42)
    n = 100
    df_X = pd.DataFrame({
        "age": np.random.randint(22, 60, size=n),
        "feat_score": np.random.randn(n),
        "feat_volume": np.random.uniform(10, 100, size=n)
    })
    s_y = pd.Series((df_X["feat_score"] > 0).astype(int), name="target")

    scout = MLScoutEngine(random_seed=42, cv_folds=3)
    ml_res = scout.run_scout(df_X, s_y)

    audit_data = {
        "audit_version": "2.2.0",
        "generated_at": "2026-09-16T12:00:00",
        "db_meta": {"target_table": "customers", "total_row_count": n, "sample_row_count": n},
        "data_health": {
            "health_score": 92,
            "total_columns": 3,
            "missing_cells_ratio": 0.0,
            "duplicate_row_count": 0,
            "pii_detected": []
        },
        "missing_summary": [],
        "numeric_profiles": [],
        "xgboost_shap_analysis": {
            "baseline_score": 0.82,
            "top_features": [{"rank": 1, "feature": "feat_score", "impact_pct": 55.0, "mean_abs_shap": 0.35, "direction": "+", "correlation": 0.65, "interpretation": "높을수록 긍정"}],
            "recommendations": {}
        },
        "ml_scout": ml_res
    }

    out_xlsx = os.path.join(tmp_path, "audit_comparative_report.xlsx")
    builder = ExcelReportBuilder()
    builder.build_report(audit_data, out_xlsx)

    assert os.path.exists(out_xlsx)
    wb = openpyxl.load_workbook(out_xlsx)
    assert "AutoML_리더보드" in wb.sheetnames

    ws = wb["AutoML_리더보드"]
    # Check Title
    assert "비교 모델(대조군) 벤치마크" in ws["A1"].value

    # Scan rows in Sheet 4
    all_values = []
    for row in ws.iter_rows(values_only=True):
        all_values.append([v for v in row if v is not None])

    text_dump = " ".join(str(item) for sublist in all_values for item in sublist)

    # Section 1 checks
    assert "AI 도입 타당성 및 비교 모델(대조군) 실측 벤치마크 총괄 요약" in text_dump
    assert "전체 통계 기준선 대비 순수 향상도 (Lift %)" in text_dump
    assert "단순 연령/군집 룰 대비 순수 향상도 (Lift %)" in text_dump

    # Section 2 checks (10 columns)
    expected_headers = [
        "순위", "모델명", "모델 역할 / 분류", "알고리즘 계열", "검증 점수",
        "전체 통계(인기도) 대비 Lift (%)", "연령/군집 룰 대비 Lift (%)",
        "학습 소요시간 (초)", "추론 복잡도", "비교 우위 근거 및 최종 채택 사유 (Verdict Rationale)"
    ]
    for h in expected_headers:
        assert h in text_dump

    # Section 3 checks (Subgroup slices)
    assert "연령별 / 군집화 계층별 대조군 룰 vs AI 챔피언 상세 실측 우위표" in text_dump
    assert "세그먼트 / 연령 군집" in text_dump
    assert "세그먼트별 순수 우위도 (Lift %)" in text_dump
