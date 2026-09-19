"""
Tests for Dongguk University (DGU) Recommendation System Deliverables.
Verifies that:
1. 'dgu_recommendation_feature_journey.xlsx' has all 4 sheets and valid 6-model benchmark metrics.
2. 'dgu_data_landscape_and_api_wbs.xlsx' has all 4 sheets, 6 tables, 10 APIs, W1~W12 WBS, and 18.5 M/M resources.
3. 'dgu_executive_presentation.pptx' has 6 widescreen (16:9) slides with complete data pattern insights.
"""
import os
import sys
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import pytest
import openpyxl
from pptx import Presentation
from src.dgu.dgu_recsys_generator import DGURecSysDeliverablesGenerator


@pytest.fixture(scope="module")
def dgu_deliverables(tmp_path_factory):
    out_dir = str(tmp_path_factory.mktemp("dgu_dist"))
    generator = DGURecSysDeliverablesGenerator(output_dir=out_dir)
    results = generator.generate_all()
    return results


def test_feature_journey_excel(dgu_deliverables):
    """Verifies that the feature journey workbook contains 4 sheets and 6-model comparative benchmarks."""
    xlsx_path = dgu_deliverables["feature_journey_excel"]
    assert os.path.exists(xlsx_path)

    wb = openpyxl.load_workbook(xlsx_path)
    sheet_names = wb.sheetnames
    assert "00_총괄_추천아키텍처_및_요약" in sheet_names
    assert "01_6대_후보모델_종합_벤치마크_평가표" in sheet_names
    assert "02_기능별_피처엔지니어링_여정" in sheet_names
    assert "03_학사규칙_및_추천의도_요건정의" in sheet_names

    # Verify Sheet 0 (3 functions)
    ws0 = wb["00_총괄_추천아키텍처_및_요약"]
    f_ids = [ws0.cell(row=r, column=1).value for r in range(5, 8)]
    assert "REC_01" in f_ids
    assert "REC_02" in f_ids
    assert "REC_03" in f_ids

    # Verify Sheet 1 (Benchmarking rows and lift)
    ws1 = wb["01_6대_후보모델_종합_벤치마크_평가표"]
    models_found = []
    has_champ_lift = False
    for row in ws1.iter_rows(values_only=True):
        if row and row[1] and any(m in str(row[1]) for m in ["단순 통계", "글로벌 인기도", "계층별 인기도", "룰 베이스", "데모그래픽", "하이브리드 ML"]):
            models_found.append(row[1])
            if "하이브리드 ML" in str(row[1]) and any(c and "+" in str(c) for c in row):
                has_champ_lift = True
    assert len(models_found) >= 18  # 3 functions * 6 models = 18 rows
    assert has_champ_lift


def test_data_landscape_and_wbs_excel(dgu_deliverables):
    """Verifies that the data landscape and WBS workbook contains 4 sheets, 6 tables, 10 APIs, and 18.5 M/M."""
    xlsx_path = dgu_deliverables["data_landscape_wbs_excel"]
    assert os.path.exists(xlsx_path)

    wb = openpyxl.load_workbook(xlsx_path)
    sheet_names = wb.sheetnames
    assert "00_데이터_랜드스케이프_전체현황" in sheet_names
    assert "01_API_개발_진행상황_최신화" in sheet_names
    assert "02_주단위_WBS_및_공수산정(M_M)" in sheet_names
    assert "03_고객협의_및_추가보완계획" in sheet_names

    # Verify 6 core tables
    ws0 = wb["00_데이터_랜드스케이프_전체현황"]
    tables = [ws0.cell(row=r, column=1).value for r in range(5, 11)]
    assert "DIM_STUDENT" in tables
    assert "FACT_GPA_SEMESTER" in tables
    assert "FACT_ATTENDANCE" in tables
    assert "FACT_EXTRACURRICULAR" in tables
    assert "FACT_COUNSELING" in tables
    assert "RULE_PREREQUISITE" in tables

    # Verify 10 APIs
    ws1 = wb["01_API_개발_진행상황_최신화"]
    apis = [ws1.cell(row=r, column=1).value for r in range(5, 15)]
    assert len(apis) == 10
    assert "API_01" in apis
    assert "API_10" in apis

    # Verify WBS W1~W12 and Total M/M
    ws2 = wb["02_주단위_WBS_및_공수산정(M_M)"]
    weeks = [ws2.cell(row=r, column=2).value for r in range(6, 18)]
    assert "W1" in weeks
    assert "W12" in weeks

    # Check Total MM cell
    total_mm_found = False
    for row in ws2.iter_rows(values_only=True):
        if row and any("18.5 M/M" in str(c) for c in row if c):
            total_mm_found = True
            break
    assert total_mm_found


def test_executive_pptx(dgu_deliverables):
    """Verifies that the 16:9 executive presentation deck contains 6 slides with expected topics."""
    pptx_path = dgu_deliverables["executive_pptx"]
    assert os.path.exists(pptx_path)

    prs = Presentation(pptx_path)
    # Widescreen check (13.333 x 7.5 inches)
    assert abs(prs.slide_width.inches - 13.333) < 0.01
    assert abs(prs.slide_height.inches - 7.5) < 0.01
    assert len(prs.slides) == 6

    # Verify Slide texts
    slide_texts = []
    for s in prs.slides:
        text_content = ""
        for shape in s.shapes:
            if shape.has_text_frame:
                text_content += " " + shape.text_frame.text
        slide_texts.append(text_content)

    assert "동국대학교 맞춤형 추천 시스템" in slide_texts[0]
    assert "핵심 패턴" in slide_texts[1]
    assert "후보 모델" in slide_texts[2] or "벤치마크" in slide_texts[2]
    assert "3단계 진화 로드맵" in slide_texts[3]
    assert "우리는 이런 것도 할 수 있습니다" in slide_texts[4]
    assert "18.5 M/M" in slide_texts[5] or "실행 계획" in slide_texts[5]
